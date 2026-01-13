"""
Export generated documentation to Microsoft Word (DOCX) format.
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re


def export_to_docx(markdown_text: str, output_file: str):
    """
    Convert markdown text to a formatted Word document.
    
    Args:
        markdown_text: The markdown-formatted documentation text
        output_file: Path where the DOCX file should be saved
    """
    
    # Create a new Document
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # Split text into lines
    lines = markdown_text.split('\n')
    
    in_code_block = False
    in_table = False
    table_lines = []
    
    for line in lines:
        line_stripped = line.strip()
        
        # Skip empty lines (but add spacing)
        if not line_stripped:
            if not in_code_block and not in_table:
                doc.add_paragraph()
            continue
        
        # Handle code blocks
        if line_stripped.startswith('```'):
            in_code_block = not in_code_block
            continue
        
        if in_code_block:
            p = doc.add_paragraph(line)
            p.style = 'Normal'
            p_format = p.paragraph_format
            p_format.left_indent = Inches(0.5)
            for run in p.runs:
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
            continue
        
        # Handle tables
        if line_stripped.startswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line_stripped)
            continue
        else:
            if in_table:
                # Process accumulated table
                _add_table_to_doc(doc, table_lines)
                in_table = False
                table_lines = []
        
        # Handle headers
        if line_stripped.startswith('#'):
            level = len(line_stripped) - len(line_stripped.lstrip('#'))
            text = line_stripped.lstrip('#').strip()
            
            heading = doc.add_heading(text, level=min(level, 9))
            
            # Style the heading
            if level == 1:
                heading.runs[0].font.size = Pt(24)
                heading.runs[0].font.color.rgb = RGBColor(31, 78, 120)
            elif level == 2:
                heading.runs[0].font.size = Pt(18)
                heading.runs[0].font.color.rgb = RGBColor(31, 78, 120)
            elif level == 3:
                heading.runs[0].font.size = Pt(14)
                heading.runs[0].font.color.rgb = RGBColor(79, 129, 189)
            
            continue
        
        # Handle horizontal rules
        if line_stripped in ['---', '___', '***', '________________________________________']:
            doc.add_paragraph('_' * 80)
            continue
        
        # Handle bullet lists
        if line_stripped.startswith(('- ', '* ', '+ ')):
            text = line_stripped[2:].strip()
            text = _clean_markdown(text)
            p = doc.add_paragraph(text, style='List Bullet')
            continue
        
        # Handle numbered lists
        if re.match(r'^\d+\.\s', line_stripped):
            text = re.sub(r'^\d+\.\s', '', line_stripped).strip()
            text = _clean_markdown(text)
            p = doc.add_paragraph(text, style='List Number')
            continue
        
        # Handle regular paragraphs
        text = _clean_markdown(line_stripped)
        p = doc.add_paragraph()
        
        # Parse inline markdown (bold, italic, code)
        _add_formatted_text(p, text)
    
    # Process any remaining table
    if in_table and table_lines:
        _add_table_to_doc(doc, table_lines)
    
    # Save the document
    doc.save(output_file)
    print(f"✓ DOCX exported to: {output_file}")


def _add_table_to_doc(doc, table_lines):
    """Add a markdown table to the document."""
    if len(table_lines) < 2:
        return
    
    # Parse table rows
    rows = []
    for line in table_lines:
        if line.startswith('|---') or line.startswith('| ---'):
            continue  # Skip separator line
        
        cells = [cell.strip() for cell in line.split('|')[1:-1]]  # Remove first/last empty elements
        if cells:
            rows.append(cells)
    
    if not rows:
        return
    
    # Create table
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Light Grid Accent 1'
    
    # Fill table
    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            cell = table.rows[i].cells[j]
            cell.text = cell_text.strip()
            
            # Make header row bold
            if i == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.size = Pt(11)
    
    # Add spacing after table
    doc.add_paragraph()


def _clean_markdown(text):
    """Remove markdown formatting but keep the text."""
    # Remove image syntax
    text = re.sub(r'!\[([^\]]*)\]\([^\)]*\)', r'\1', text)
    
    # Remove link syntax but keep text
    text = re.sub(r'\[([^\]]*)\]\([^\)]*\)', r'\1', text)
    
    return text


def _add_formatted_text(paragraph, text):
    """Add text with inline markdown formatting (bold, italic, code)."""
    
    # Split by code blocks first
    parts = re.split(r'(`[^`]+`)', text)
    
    for part in parts:
        if part.startswith('`') and part.endswith('`'):
            # Code text
            run = paragraph.add_run(part.strip('`'))
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(200, 0, 0)
        else:
            # Split by bold/italic
            subparts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', part)
            
            for subpart in subparts:
                if subpart.startswith('**') and subpart.endswith('**'):
                    # Bold text
                    run = paragraph.add_run(subpart.strip('*'))
                    run.font.bold = True
                elif subpart.startswith('*') and subpart.endswith('*'):
                    # Italic text
                    run = paragraph.add_run(subpart.strip('*'))
                    run.font.italic = True
                else:
                    # Normal text
                    if subpart:
                        paragraph.add_run(subpart)